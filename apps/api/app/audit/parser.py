"""Hierarchical regulation parser over the ingestion pipeline's page records.

Ordinary DOCX/PDF/TXT clauses retain their numbered/lettered hierarchy.
Explicit DOCX/XLSX table headers define sourced unit/role rows and duty cells;
unheaded single-column numbered annexes retain their individual cell locations.
Locations refer to physical PDF pages, DOCX body blocks or Excel cells, never
to an invented Word page number.
"""

from __future__ import annotations

import re
from dataclasses import dataclass

from .models import Citation, Clause, Document, ParseResult, SourceLocation, Unit
from .text import is_governance_statement, normalize

_NUM_START = re.compile(r"^(?P<label>(?P<path>\d{1,3}(?:\.\d{1,3})*)\.)(?=\s|[^\d\s.])")
_LETTER_START = re.compile(r"^(?P<label>(?P<letter>[а-яё])[.)])(?:\s+|(?=[А-ЯЁ]))")
# A numeric marker inside a line: after sentence-ending punctuation, followed by a capitalised word.
_NUM_EMBEDDED = re.compile(
    r"(?<=[.;:!?)»])\s+(?P<label>(?P<path>\d{1,3}(?:\.\d{1,3})*)\.)\s*(?=[А-ЯЁA-Z])"
)
# "п. 3." / "No 7." style cross-references are not clause markers.
_REF_BEFORE = re.compile(
    r"(?<!\w)(?:п|пп|ст|ч|гл|разд|пункт\w*|подпункт\w*|раздел\w*|№|No)\.?\s*$", re.IGNORECASE
)
_TOC_TITLE = re.compile(r"^(оглавление|содержание)\s*$", re.IGNORECASE)
_TOC_ENTRY_TAIL = re.compile(r"\s\d{1,4}\s*$")

# A list-introducing clause about composition or subordination is org structure.
_COMPOSITION = re.compile(r"состо[ия]т\s+из|входят\s+следующ|в\s+состав\s+\S+\s+входят", re.IGNORECASE)
_SUBORDINATION = re.compile(r"подчиня[ею]тся|в\s+составе\s+следующих\s+должностей", re.IGNORECASE)

# Depth-1 clauses are section headings unless their block carries substantive body text.
_HEADING_BODY_LIMIT = 200

_UNIT_WORDS = r"(?:Блок|Департамент|Управлени|Отдел|Служб|Центр|Дирекци|Комитет|Сектор|Групп|Подразделени)"
_ABBR_DEF = re.compile(
    r"(?P<name>" + _UNIT_WORDS + r"[^():;,]{0,80}?)\s*\((?:далее\s*[-–—]?\s*)?(?P<abbr>[А-ЯЁA-Z][^()]{0,40}?)\)"
)
_WORKERS_OF = re.compile(r"работник\w*\s+(?P<abbr>[А-ЯЁA-Z][А-ЯЁA-Z-]+)\b")

# A governing role is a heading/list definition, not a role merely mentioned in
# a duty. Keep the title separate from the (edition-dependent) description.
_ROLE_NOUN = (
    r"(?:директор\w*|руководител\w*|начальник\w*|заместител\w*|"
    r"менеджер\w*|аудитор\w*|инженер\w*|юрист\w*|бухгалтер\w*|"
    r"работник\w*|председател\w*|специалист\w*|координатор\w*|эксперт\w*)"
)
_ROLE_START = re.compile(
    rf"^(?:(?:[а-яё]+(?:ый|ий|ой|ая|ого|его|ому|ему)\s+){{0,2}}{_ROLE_NOUN})\b",
    re.IGNORECASE,
)
_ROLE_PREDICATE = re.compile(
    r"\s+(?:(?:не\s+)?(?:обязан\w*|должен|должна|должны|"
    r"нес\w*\s+ответственность\s+за|отвеча\w*\s+за|"
    r"име\w*\s+право|име\w*\s+права|вправе|может|могут)\b|"
    r"запрещен\w*\b)",
    re.IGNORECASE,
)
_ROLE_LABEL = re.compile(r"^(?:функции|обязанности|ответственность|полномочия)\s+", re.IGNORECASE)
_ROLE_ALIAS = re.compile(r"\s*\(далее\b[^()]*\)$", re.IGNORECASE)
_ROLE_ACTION = re.compile(
    r"\b(?:осуществля\w*|организу\w*|обеспечива\w*|"
    r"поруча\w*|поручен\w*|делегир\w*|возлага\w*|переда\w*|"
    r"согласовыва\w*|утвержда\w*|представля\w*|"
    r"нес\w*|обязан\w*|долж\w*|име\w*|отвеча\w*|"
    r"вправе|может|могут|запрещен\w*)\b",
    re.IGNORECASE,
)
_ROLE_FINITE_VERB = re.compile(
    r"\b[а-яё]{4,}(?:ется|ются|ится|атся|яются|ается|"
    r"ет|ют|ут|ит|ят|ают|яют)\b", re.IGNORECASE,
)

_TABLE_NUMBER = re.compile(r"(?:№|номер|п п|number|no)(?:\s+(?:пункта|строки))?$")


@dataclass
class _Segment:
    label: str
    kind: str  # "num" | "letter" | "plain" | "toc" | "table_*"
    path: tuple[int, ...] | None
    letter: str | None
    text: str
    location: SourceLocation | None = None
    row_key: tuple[int, int] | None = None
    anchor_key: tuple[int, int] | None = None
    parent_name: str = ""
    parent_key: tuple[int, int] | None = None
    anchor_kind: str | None = None
    unit_reference: bool = False


@dataclass
class _TableDefinition:
    clause: Clause
    kind: str
    parent_name: str
    row_key: tuple[int, int]
    reference: bool = False

def _parse_path(raw: str) -> tuple[int, ...]:
    return tuple(int(part) for part in raw.split("."))


def _is_successor(candidate: tuple[int, ...], current: tuple[int, ...] | None) -> bool:
    """True when `candidate` is the first child or a next sibling at some level of `current`."""
    if current is None:
        return candidate == (1,)
    if candidate == current + (1,):
        return True
    for level in range(len(current)):
        if candidate == current[:level] + (current[level] + 1,):
            return True
    return False


def _split_embedded(label: str, kind: str, path, letter, body: str, current):
    """Split `body` at embedded numeric markers that continue the numbering."""
    segments: list[_Segment] = []
    start = 0
    head_label, head_kind, head_path, head_letter = label, kind, path, letter
    running = path if path is not None else current
    for match in _NUM_EMBEDDED.finditer(body):
        candidate = _parse_path(match.group("path"))
        if not _is_successor(candidate, running) or _REF_BEFORE.search(body[: match.start()]):
            continue
        segments.append(_Segment(head_label, head_kind, head_path, head_letter, body[start : match.start()].strip()))
        head_label, head_kind, head_path, head_letter = match.group("label"), "num", candidate, None
        running = candidate
        start = match.end()
    segments.append(_Segment(head_label, head_kind, head_path, head_letter, body[start:].strip()))
    return segments, running


def _source_location(page: dict, row: dict | None = None, column: int | None = None) -> SourceLocation | None:
    if "sheet" in page:
        from openpyxl.utils import get_column_letter

        if row is None:
            return SourceLocation(sheet=page["sheet"])
        number = row["row"]
        if column is not None:
            cell_range = f"{get_column_letter(column)}{number}"
            for span in page.get("merges", ()):
                if span["min_row"] == number and span["min_col"] == column:
                    if span["max_row"] != number or span["max_col"] != column:
                        cell_range += f":{get_column_letter(span['max_col'])}{span['max_row']}"
                    break
        else:
            used = [index for index, value in enumerate(row["cells"], start=1) if value.strip()]
            cell_range = (
                f"{get_column_letter(used[0])}{number}:{get_column_letter(used[-1])}{number}"
                if len(used) > 1 else f"{get_column_letter(used[0])}{number}"
            )
        return SourceLocation(sheet=page["sheet"], cell_range=cell_range)
    if "block" in page:
        return SourceLocation(block=page["block"])
    if "physical_page" in page:
        return SourceLocation(page=page["physical_page"])
    return None


def _header_kind(cell: str) -> str | None:
    """Recognise explicit column labels, not arbitrary duty or unit names."""
    if cell.strip() == "№":
        return "number"
    text = normalize(cell)
    if not text or len(text) > 70:
        return None
    if _TABLE_NUMBER.fullmatch(text):
        return "number"
    if re.fullmatch(
        r"(?:родительск\w*|вышестоящ\w*|головн\w*|"
        r"кому подчиняется|в составе|parent|reports to)"
        r"(?:\s+(?:структурн\w*\s+)?подразделени\w*|\s+unit)?", text,
    ):
        return "parent"
    stem = re.sub(r"^(?:наименование|название|описание|перечень|основные|виды)\s+", "", text)
    if re.fullmatch(
        r"(?:структурн\w*\s+)?(?:подразделени\w*|департамент\w*|"
        r"управлени\w*|отдел\w*|служб\w*|центр\w*|unit|department|division)", stem,
    ):
        return "unit"
    if re.fullmatch(
        r"(?:должност\w*|рол\w*|ответственн\w*\s+исполнител\w*|"
        r"исполнител\w*|role|position)", stem,
    ):
        return "role"
    if re.fullmatch(
        r"(?:функциональн\w*\s+)?(?:функци\w*|обязанност\w*|задач\w*|полномочи\w*|"
        r"ответственност\w*|действи\w*|работ\w*|"
        r"function\w*|dut\w*|responsibilit\w*|task\w*)"
        r"(?:\s+и\s+(?:функци\w*|обязанност\w*|задач\w*))?"
        r"(?:\s+(?:структурн\w*\s+)?подразделени\w*)?", stem,
    ):
        return "function"
    return None


def _table_header(cells: list[str]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for column, cell in enumerate(cells, start=1):
        kind = _header_kind(cell)
        if kind is not None:
            if kind in columns:
                return {}  # Ambiguous duplicate columns must not imply ownership.
            columns[kind] = column
    return columns if any(kind in columns for kind in ("unit", "role", "function")) else {}


def _merged_anchor(page: dict, row: dict, column: int) -> tuple[int, int] | None:
    if column in row.get("merged", {}):
        return tuple(row["merged"][column])
    for span in page.get("merges", ()):
        if (span["min_row"] <= row["row"] <= span["max_row"]
                and span["min_col"] <= column <= span["max_col"]
                and (row["row"], column) != (span["min_row"], span["min_col"])):
            return span["min_row"], span["min_col"]
    return None


def _table_segments(page: dict) -> tuple[list[_Segment], bool]:
    """Convert explicitly headed table rows into individual source cells."""
    segments: list[_Segment] = []
    columns: dict[str, int] = {}
    recognised = False
    prior_row: int | None = None
    header_row = 0
    rows_by_number = {row["row"]: row for row in page["rows"]}
    for row in page["rows"]:
        if prior_row is not None and row["row"] > prior_row + 1:
            columns = {}  # A blank intervening row ends a table.
        prior_row = row["row"]
        header = _table_header(row["cells"])
        if header and (not columns or header == columns or len(header) >= 2):
            columns = header
            recognised = True
            header_row = row["row"]
            segments.append(_Segment("", "table_header", None, None, row["text"],
                                     _source_location(page, row), (page["page"], row["row"])))
            continue
        if not columns:
            segments.append(_Segment("", "table_other", None, None, row["text"],
                                     _source_location(page, row), (page["page"], row["row"])))
            continue
        row_key = (page["page"], row["row"])
        typed_columns = {columns[k] for k in ("unit", "role", "function", "parent") if k in columns}
        for column, value in enumerate(row["cells"], start=1):
            if value.strip() and (column not in typed_columns or column in row.get("formulas", ())):
                segments.append(_Segment(
                    "", "table_other", None, None, value,
                    _source_location(page, row, column), row_key,
                ))
        parent_col = columns.get("parent")
        parent_name = (
            row["cells"][parent_col - 1].strip()
            if parent_col is not None and parent_col <= len(row["cells"])
            and parent_col not in row.get("formulas", ()) else ""
        )
        parent_anchor = _merged_anchor(page, row, parent_col) if parent_col else None
        parent_key = None
        if parent_anchor and parent_anchor[0] > header_row and not parent_name:
            anchor_row = rows_by_number.get(parent_anchor[0])
            if (anchor_row and len(anchor_row["cells"]) >= parent_anchor[1]
                    and parent_anchor[1] not in anchor_row.get("formulas", ())):
                parent_name = anchor_row["cells"][parent_anchor[1] - 1].strip()
                if parent_name:
                    parent_key = (page["page"], parent_anchor[0])
        label_col = columns.get("number")
        label = (
            row["cells"][label_col - 1]
            if label_col is not None and label_col <= len(row["cells"])
            and label_col not in row.get("formulas", ()) else ""
        )
        if parent_name and parent_key is None:
            segments.append(_Segment(
                "", "table_parent", None, None, row["cells"][parent_col - 1],
                _source_location(page, row, parent_col), row_key,
            ))
        for kind in ("unit", "role", "function"):
            column = columns.get(kind)
            if column is None:
                continue
            value = row["cells"][column - 1] if column <= len(row["cells"]) else ""
            if not value.strip():
                continue
            if column in row.get("formulas", ()):
                continue
            owner_anchor = None
            owner_kind_anchor = None
            if kind == "function":
                for owner_kind in ("role", "unit"):
                    owner_col = columns.get(owner_kind)
                    if owner_col is not None:
                        merged = _merged_anchor(page, row, owner_col)
                        if merged and merged[0] > header_row:
                            owner_anchor = (page["page"], merged[0])
                            owner_kind_anchor = owner_kind
                            break
            elif kind == "role":
                unit_col = columns.get("unit")
                if unit_col is not None:
                    merged = _merged_anchor(page, row, unit_col)
                    if merged and merged[0] > header_row:
                        owner_anchor = (page["page"], merged[0])
                        owner_kind_anchor = "unit"
            segments.append(_Segment(
                label if kind == "function" else "", f"table_{kind}", None, None, value,
                _source_location(page, row, column), row_key,
                owner_anchor, parent_name, parent_key, owner_kind_anchor,
                kind == "unit" and ("function" in columns or "role" in columns),
            ))
        if not any(seg.row_key == row_key and seg.kind.startswith("table_") for seg in segments[-4:]):
            segments.append(_Segment("", "table_other", None, None, row["text"],
                                     _source_location(page, row), row_key))
    return segments, recognised


def _flow_column(page: dict) -> int | None:
    """Only a genuinely one-column, sequentially numbered sheet is flowing text."""
    columns = {
        column
        for row in page["rows"]
        for column, value in enumerate(row["cells"], start=1)
        if value.strip()
    }
    if len(columns) != 1 or any(
        span["min_col"] != span["max_col"] for span in page.get("merges", ())
    ):
        return None
    column = next(iter(columns))
    paths = [
        _parse_path(match.group("path"))
        for row in page["rows"]
        if column not in row.get("formulas", ())
        for line in row["cells"][column - 1].splitlines()
        if (match := _NUM_START.match(line.strip().lstrip("\ufeff").strip()))
    ]
    if len(paths) < 2 or any(
        not _is_successor(later, earlier)
        for earlier, later in zip(paths, paths[1:])
    ):
        return None
    return column


def _segments(pages: list[dict]) -> tuple[list[_Segment], int, list[str]]:
    """Turn source pages and table rows into labelled, located segments."""
    out: list[_Segment] = []
    current: tuple[int, ...] | None = None
    in_toc = False
    embedded = 0
    warnings: list[str] = []

    def append_lines(page: dict, text: str, location: SourceLocation | None) -> None:
        nonlocal current, in_toc, embedded
        can_continue = False
        previous_break = ""
        for raw_line in text.splitlines(keepends=True):
            line = raw_line.strip().lstrip("\ufeff").strip()
            line_break = raw_line[len(raw_line.rstrip("\r\n")):]
            if not line or page.get("extraction") == "unreadable":
                can_continue = False
                previous_break = line_break
                continue
            num = _NUM_START.match(line)
            letter = _LETTER_START.match(line) if not num else None
            # PDF extraction can wrap a paragraph at a physical line boundary.
            # Attach only a lowercase continuation to its preceding labelled
            # clause on this page; do not infer order across cells or columns.
            if (
                "physical_page" in page and can_continue and not num and not letter
                and not _NUM_EMBEDDED.search(line)
                and re.match(r"^[a-zа-яё]", line)
            ):
                out[-1].text += previous_break + line
                can_continue = line[-1] not in ".!?:;"
                previous_break = line_break
                continue
            if _TOC_TITLE.match(line):
                in_toc = True
                out.append(_Segment("", "plain", None, None, line, location))
                can_continue = False
                previous_break = line_break
                continue
            if in_toc:
                if num and _TOC_ENTRY_TAIL.search(line):
                    out.append(_Segment(num.group("label"), "toc", _parse_path(num.group("path")), None,
                                        line[num.end():].strip(), location))
                    can_continue = False
                    previous_break = line_break
                    continue
                in_toc = False
            if num:
                path = _parse_path(num.group("path"))
                segs, current = _split_embedded(num.group("label"), "num", path, None, line[num.end():], current)
            elif letter:
                segs, current = _split_embedded(letter.group("label"), "letter", None, letter.group("letter"),
                                                line[letter.end():], current)
            else:
                segs, current = _split_embedded("", "plain", None, None, line, current)
            embedded += len(segs) - 1
            for seg in segs:
                seg.location = location
            out.extend(segs)
            can_continue = (
                "physical_page" in page and segs[-1].kind in ("num", "letter")
                and line[-1] not in ".!?:;"
            )
            previous_break = line_break

    for page in pages:
        if "rows" in page:
            table_segments, recognised = _table_segments(page)
            formula_count = sum(len(row.get("formulas", ())) for row in page["rows"])
            if formula_count:
                warnings.append(
                    f"Sheet {page['sheet']!r}: {formula_count} formula cell(s) retained as "
                    "source text but not evaluated as units, parents or duties."
                )
            column = None if recognised else _flow_column(page)
            if column is not None:
                current = None
                in_toc = False
                for row in page["rows"]:
                    value = row["cells"][column - 1]
                    location = _source_location(page, row, column)
                    if column in row.get("formulas", ()):
                        out.append(_Segment("", "table_other", None, None, value, location,
                                            (page["page"], row["row"])))
                    else:
                        append_lines(page, value, location)
                continue
            out.extend(table_segments)
            if "sheet" in page and page["rows"] and not recognised:
                warnings.append(
                    f"Sheet {page['sheet']!r}: no explicit unit/function column headers; "
                    "rows retained as source only, not invented duties."
                )
            continue
        append_lines(page, str(page.get("text") or ""), _source_location(page))
    return out, embedded, warnings


def _classify_numeric(path: tuple[int, ...], text: str) -> str:
    if not any(char.isalnum() for char in text):
        return "other"
    if len(path) == 1:
        return "function" if len(text) > _HEADING_BODY_LIMIT else "heading"
    if text.rstrip().endswith(":") and (_COMPOSITION.search(text) or _SUBORDINATION.search(text)):
        return "structure"
    return "function"


def _parse_clauses(
    doc: str, pages: list[dict],
) -> tuple[list[Clause], list[str], list[_TableDefinition]]:
    segments, embedded, warnings = _segments(pages)
    clauses: list[Clause] = []
    table_definitions: list[_TableDefinition] = []
    table_rows: dict[tuple[int, int], dict[str, str]] = {}
    used: dict[str, int] = {}
    latest_by_path: dict[tuple[int, ...], str] = {}
    kind_by_id: dict[str, str] = {}
    numeric_ctx: str | None = None  # latest numeric clause id
    labelled_ctx: str | None = None  # latest numeric or letter clause id
    toc_parent: str | None = None
    orphan_letters = 0
    unlabelled = 0
    duplicates: list[str] = []

    def unique(base: str) -> str:
        count = used.get(base, 0) + 1
        used[base] = count
        if count == 1:
            return base
        duplicates.append(f"{base}@{count}")
        return f"{base}@{count}"

    for ordinal, seg in enumerate(segments, start=1):
        if seg.kind == "num":
            base = ".".join(str(p) for p in seg.path)
            clause_id = unique(base)
            parent_id = latest_by_path.get(seg.path[:-1]) if len(seg.path) > 1 else None
            kind = _classify_numeric(seg.path, seg.text)
            for key in [k for k in latest_by_path if len(k) > len(seg.path)]:
                del latest_by_path[key]
            latest_by_path[seg.path] = clause_id
            numeric_ctx = labelled_ctx = clause_id
            toc_parent = None
        elif seg.kind == "letter" and numeric_ctx is not None:
            clause_id = unique(f"{numeric_ctx}/{seg.letter}")
            parent_id = numeric_ctx
            kind = "structure" if kind_by_id.get(numeric_ctx) == "structure" else "function"
            labelled_ctx = clause_id
        elif seg.kind == "toc":
            base = ".".join(str(p) for p in seg.path)
            clause_id = unique(base)
            parent_id = toc_parent
            kind = "other"
        elif seg.kind.startswith("table_"):
            assert seg.row_key is not None
            page_number, row_number = seg.row_key
            tag = {"table_unit": "u", "table_role": "o", "table_function": "f",
                   "table_parent": "p", "table_header": "h", "table_other": "x"}[seg.kind]
            clause_id = unique(f"@t{page_number}r{row_number}{tag}")
            related = table_rows.get(seg.row_key, {})
            anchored = table_rows.get(seg.anchor_key, {}) if seg.anchor_key else {}
            if seg.kind == "table_function":
                parent_id = (
                    related.get("role")
                    or (anchored.get("role") if seg.anchor_kind == "role" else None)
                    or related.get("unit")
                    or (anchored.get("unit") if seg.anchor_kind == "unit" else None)
                )
                kind = "function"
            elif seg.kind in ("table_unit", "table_role"):
                if seg.kind == "table_role":
                    parent_id = (
                        related.get("unit") or (anchored.get("unit") if seg.anchor_kind == "unit" else None) or related.get("parent")
                        or table_rows.get(seg.parent_key, {}).get("parent")
                    )
                else:
                    parent_id = (
                        related.get("parent")
                        or table_rows.get(seg.parent_key, {}).get("parent")
                    )
                kind = "structure"
                table_rows.setdefault(seg.row_key, {})[seg.kind.removeprefix("table_")] = clause_id
            else:
                parent_id = labelled_ctx
                kind = "other"
                if seg.kind == "table_parent":
                    table_rows.setdefault(seg.row_key, {})["parent"] = clause_id
        else:
            if seg.kind == "letter":
                orphan_letters += 1
            else:
                unlabelled += 1
            clause_id = f"@p{ordinal}"
            used[clause_id] = 1
            parent_id = labelled_ctx
            kind = "other"
            if _TOC_TITLE.match(seg.text):
                toc_parent = clause_id
                kind = "heading"
        kind_by_id[clause_id] = kind
        clause = Clause(
            doc=doc, clause_id=clause_id, label=seg.label, parent_id=parent_id,
            text=seg.text, ordinal=ordinal, kind=kind, unit_ids=[], location=seg.location,
        )
        clauses.append(clause)
        if seg.kind in ("table_unit", "table_role"):
            table_definitions.append(_TableDefinition(
                clause, seg.kind.removeprefix("table_"), seg.parent_name, seg.row_key, seg.unit_reference,
            ))

    if not any(c.kind == "function" for c in clauses):
        warnings.append(
            f"{doc}: no numbered or explicitly headed function clauses were recognised; "
            "function comparison is unavailable for this document."
        )
    for page in pages:
        if page.get("extraction") == "unreadable":
            warnings.append(
                f"{doc}: physical PDF page {page['physical_page']} has no extractable text; "
                "image/diagram content was not interpreted."
            )
        elif page.get("extraction") == "ocr":
            warnings.append(
                f"{doc}: physical PDF page {page['physical_page']} used OCR text; "
                "verify recognition manually; diagram relations are not interpreted."
            )
        elif page.get("extraction") == "limited":
            warnings.append(
                f"{doc}: physical PDF page {page['physical_page']} has very little extractable "
                "text; image/diagram content may be missing and was not interpreted."
            )
    if embedded:
        warnings.append(f"{doc}: {embedded} numbered clause(s) were split out of a shared text block.")
    if unlabelled:
        warnings.append(f"{doc}: {unlabelled} unlabelled block(s) kept as kind=other and excluded from alignment.")
    empty_numbered = [
        c.clause_id for c in clauses
        if c.label and c.kind == "other" and not any(char.isalnum() for char in c.text)
    ]
    if empty_numbered:
        shown = ", ".join(empty_numbered[:10])
        warnings.append(
            f"{doc}: {len(empty_numbered)} numbered marker(s) have no function text "
            f"and were excluded from alignment: {shown}."
        )
    if orphan_letters:
        warnings.append(f"{doc}: {orphan_letters} lettered block(s) had no numbered parent; kept as kind=other.")
    if duplicates:
        shown = ", ".join(duplicates[:10]) + (" …" if len(duplicates) > 10 else "")
        warnings.append(f"{doc}: repeated clause numbers disambiguated as {shown}.")
    return clauses, warnings, table_definitions


def parse_clauses(doc: str, pages: list[dict]) -> tuple[list[Clause], list[str]]:
    """Parse one document into unique, located clauses; preserve the public API."""
    clauses, warnings, _ = _parse_clauses(doc, pages)
    return clauses, warnings


def _quote_name(text: str) -> str:
    return text.strip().rstrip(".;, ").strip()

def _role_names(text: str, *, definition: bool = False) -> list[str]:
    """Return literal titles only from a role header or enumerated definition."""
    if definition:
        title = re.split(r"\s+[—–-]\s+", _quote_name(text), maxsplit=1)[0]
    else:
        body = text.strip()
        if not body.endswith(":") or len(body) > 400:
            return []
        body = _ROLE_LABEL.sub("", body[:-1].strip())
        predicate = _ROLE_PREDICATE.search(body)
        title = body[:predicate.start()].strip() if predicate else body
        title = _ROLE_ALIAS.sub("", title).strip()
        if (re.search(r"[;,.!?]", title) or _ROLE_FINITE_VERB.search(title)
                or re.search(r"\b(?:не|должен|должна|должны|может|могут|вправе)\b", title, re.IGNORECASE)):
            return []
    names: list[str] = []
    start = 0
    for conjunction in re.finditer(r"\s+и\s+", title, re.IGNORECASE):
        if _ROLE_START.match(title[conjunction.end():]):
            names.append(title[start:conjunction.start()].strip())
            start = conjunction.end()
    names.append(title[start:].strip())
    if not names or any(
        not _ROLE_START.match(name) or _ROLE_ACTION.search(name) or len(name) > 100
        for name in names
    ):
        return []
    return names



def extract_units(
    doc: str, clauses: list[Clause], *, table_definitions: list[_TableDefinition] | None = None,
) -> list[Unit]:
    """Extract source-backed units and roles and attach their associations."""
    units: list[Unit] = []
    table_definitions = table_definitions or []
    table_clause_ids = {definition.clause.clause_id for definition in table_definitions}
    table_clause_ids.update(c.clause_id for c in clauses if c.clause_id.startswith("@t"))
    by_key: dict[str, Unit] = {}
    ids: set[str] = set()
    children: dict[str, list[Clause]] = {}
    for clause in clauses:
        if clause.parent_id is not None:
            children.setdefault(clause.parent_id, []).append(clause)

    def new_id(clause_id: str) -> str:
        unit_id, n = clause_id, 1
        while unit_id in ids:
            n += 1
            unit_id = f"{clause_id}#{n}"
        ids.add(unit_id)
        return unit_id

    def add(clause: Clause, name: str, kind: str, key: str, parent_unit_id: str | None,
            quote: str | None = None) -> Unit:
        unit = Unit(doc=doc, unit_id=new_id(clause.clause_id), name=name, kind=kind,
                    parent_unit_id=parent_unit_id,
                    citations=[Citation(doc=doc, clause_id=clause.clause_id, quote=quote or name)])
        units.append(unit)
        if key and key not in by_key:
            by_key[key] = unit
        return unit

    structure_ids = {c.clause_id for c in clauses if c.kind == "structure"}
    # Pass 1: abbreviation definitions outside structure lists.
    for clause in clauses:
        if (clause.kind == "structure" or clause.parent_id in structure_ids
                or clause.clause_id in table_clause_ids):
            continue
        for match in _ABBR_DEF.finditer(clause.text):
            abbr = match.group("abbr").strip()
            if abbr in by_key:
                continue
            name = match.group("name").strip()
            add(clause, f"{name} ({abbr})", "unit", abbr, None, quote=match.group(0).strip())

    # Pass 2: enumerated composition / subordination lists.
    for clause in clauses:
        if clause.kind != "structure" or clause.clause_id in table_clause_ids:
            continue
        is_roles = bool(_SUBORDINATION.search(clause.text)) and not _COMPOSITION.search(clause.text)
        parent_unit: Unit | None = None
        if is_roles:
            owner = _WORKERS_OF.search(clause.text)
            if owner:
                parent_unit = by_key.get(owner.group("abbr"))
        else:
            head = _COMPOSITION.split(clause.text, maxsplit=1)[0]
            parent_unit = next((u for k, u in by_key.items() if re.search(rf"(?<!\w){re.escape(k)}(?!\w)", head)), None)
        for child in children.get(clause.clause_id, []):
            if child.kind != "structure" or child.clause_id in table_clause_ids:
                continue
            name = _quote_name(child.text)
            if not name:
                continue
            if is_roles:
                role_names = _role_names(child.text, definition=True) or [name]
                for role_name in role_names:
                    unit = add(child, role_name, "role", "", parent_unit.unit_id if parent_unit else None,
                               quote=role_name)
                    child.unit_ids.append(unit.unit_id)
                continue
            abbr_match = re.search(r"\(([А-ЯЁA-Z][А-ЯЁA-Z-]+)\)", name)
            key = abbr_match.group(1) if abbr_match else ""
            existing = by_key.get(key) if key else None
            if existing is not None and existing.kind == "unit":
                units.remove(existing)
                ids.discard(existing.unit_id)
                del by_key[key]
            unit = add(child, name, "unit", key, parent_unit.unit_id if parent_unit else None)
            child.unit_ids = [unit.unit_id]

    # Duties after a header do not become part of the role's identity.
    for clause in clauses:
        if (clause.kind == "structure" or clause.parent_id in structure_ids
                or clause.clause_id in table_clause_ids):
            continue
        for role_name in _role_names(clause.text):
            unit = add(clause, role_name, "role", "", None, quote=role_name)
            clause.unit_ids.append(unit.unit_id)

    # Load structural tables first. A duty-table unit cell may refer to one
    # uniquely named structural unit elsewhere in this same document.
    table_units: dict[tuple[int, str, str, str], Unit] = {}
    row_units: dict[tuple[int, int], dict[str, Unit]] = {}
    table_units_by_clause: dict[str, Unit] = {}
    declarations = [d for d in table_definitions if d.kind == "unit" and not d.reference]
    references = [d for d in table_definitions if d.kind != "unit" or d.reference]
    declared_units: dict[tuple[str, str], list[Unit]] = {}
    for definition in declarations + references:
        clause = definition.clause
        name = clause.text.strip()
        if not name:
            continue
        parent = table_units_by_clause.get(clause.parent_id) if definition.kind == "role" else None
        if definition.kind == "role":
            owner = (
                f"unit:{parent.unit_id}" if parent and parent.kind == "unit"
                else f"name:{normalize(definition.parent_name)}" if definition.parent_name
                else f"unscoped:{clause.clause_id}"
            )
        else:
            owner = normalize(definition.parent_name)
        key = (definition.row_key[0], definition.kind, normalize(name), owner)
        unit = table_units.get(key)
        if unit is None and definition.kind == "unit" and definition.reference:
            matches = declared_units.get((normalize(name), owner), [])
            if len(matches) == 1:
                unit = matches[0]
        if unit is None:
            unit = add(clause, name, definition.kind, "",
                       parent.unit_id if parent and parent.kind == "unit" else None)
            if definition.kind == "unit" and not definition.reference:
                declared_units.setdefault((normalize(name), owner), []).append(unit)
        else:
            unit.citations.append(Citation(doc=doc, clause_id=clause.clause_id, quote=name))
        table_units[key] = unit
        clause.unit_ids.append(unit.unit_id)
        row_units.setdefault(definition.row_key, {})[definition.kind] = unit
        table_units_by_clause[clause.clause_id] = unit

    for definition in table_definitions:
        if definition.kind != "role":
            continue
        role = row_units[definition.row_key]["role"]
        parent = table_units_by_clause.get(definition.clause.parent_id)
        if parent and parent.kind == "unit":
            role.parent_unit_id = parent.unit_id
    # Resolve only unambiguous source-cell parent names to actual units.
    by_clause = {clause.clause_id: clause for clause in clauses}
    for definition in table_definitions:
        if definition.kind not in ("unit", "role") or not definition.parent_name:
            continue
        child = row_units[definition.row_key][definition.kind]
        if child.parent_unit_id is not None:
            continue
        scope = definition.row_key[0]
        candidates = [
            parent for (page, kind, name, _), parent in table_units.items()
            if page == scope and kind == "unit" and parent.unit_id != child.unit_id
            and (name == normalize(definition.parent_name)
                 or normalize(unit_key(parent)) == normalize(definition.parent_name))
        ]
        if len(candidates) != 1:
            continue
        child.parent_unit_id = candidates[0].unit_id
        parent_clause = by_clause.get(definition.clause.parent_id)
        if parent_clause is not None and parent_clause.text.strip() == definition.parent_name:
            citation = Citation(doc=doc, clause_id=parent_clause.clause_id, quote=definition.parent_name)
            if citation not in child.citations:
                child.citations.append(citation)
    _attach_unit_ids(clauses, units)
    return units


def unit_key(unit: Unit) -> str:
    """Cross-document identity from a stable name, never a changing header duty."""
    if unit.kind == "role":
        return normalize(unit.name)
    match = re.search(r"\(([А-ЯЁA-Z][А-ЯЁA-Z-]+)\)\s*$", unit.name)
    if match:
        return match.group(1)
    return normalize(unit.name)


def _unit_scope(clause: Clause, unit: Unit) -> bool:
    if unit.kind != "unit" or clause.kind not in ("heading", "structure"):
        return False
    if any(c.clause_id == clause.clause_id and c.quote in clause.text for c in unit.citations):
        return True
    if normalize(_quote_name(clause.text).rstrip(":")) == normalize(unit.name):
        return True
    token = unit_key(unit)
    if not re.fullmatch(r"[А-ЯЁA-Z][А-ЯЁA-Z-]+", token):
        return False
    return bool(re.match(
        rf"^(?:(?:функции|задачи|обязанности|полномочия|ответственность)\s+)?{re.escape(token)}\b",
        clause.text.strip(), re.IGNORECASE,
    ))


def _sibling_role_scopes(
    clauses: list[Clause], direct: dict[str, list[str]], units_by_id: dict[str, Unit],
) -> dict[str, list[str]]:
    """Associate a standalone role title with following siblings in its section.

    Source stays the original @p clause; no child parent_id or kind is changed.
    A subsequent numbered role header closes that title's interval.
    """
    siblings: dict[str, list[Clause]] = {}
    # The next labelled clause establishes the list's level. A plain heading
    # may follow a deeply nested item; its raw parent is that previous item.
    next_parent: dict[str, str | None] = {}
    upcoming: Clause | None = None
    for clause in sorted(clauses, key=lambda c: c.ordinal, reverse=True):
        next_parent[clause.clause_id] = upcoming.parent_id if upcoming else None
        if clause.label:
            upcoming = clause
    for clause in clauses:
        scope_parent = clause.parent_id
        if not clause.label and clause.text.rstrip().endswith(":"):
            scope_parent = next_parent[clause.clause_id]
        if scope_parent:
            siblings.setdefault(scope_parent, []).append(clause)
    scopes: dict[str, list[str]] = {}
    for group in siblings.values():
        active: list[str] = []
        for clause in sorted(group, key=lambda c: c.ordinal):
            roles = [
                uid for uid in direct.get(clause.clause_id, [])
                if uid in units_by_id and units_by_id[uid].kind == "role"
                and any(c.clause_id == clause.clause_id and c.quote in clause.text
                        for c in units_by_id[uid].citations)
            ]
            if roles:
                active = roles if not clause.label and clause.text.rstrip().endswith(":") else []
            elif not clause.label and clause.text.rstrip().endswith(":"):
                active = []
            if clause.label and active:
                scopes[clause.clause_id] = active.copy()
    return scopes


def _scoped_sibling_roles(
    clause: Clause, clauses_by_id: dict[str, Clause], scopes: dict[str, list[str]],
) -> list[str]:
    current: Clause | None = clause
    while current is not None:
        if current.clause_id in scopes:
            return scopes[current.clause_id]
        current = clauses_by_id.get(current.parent_id) if current.parent_id else None
    return []


def _attach_unit_ids(clauses: list[Clause], units: list[Unit]) -> None:
    """Keep direct mentions as associations; inherit only a source-backed scope."""
    patterns: list[tuple[str, re.Pattern]] = []
    for unit in units:
        if unit.kind != "unit":
            continue
        token = unit_key(unit)
        if re.fullmatch(r"[А-ЯЁA-Z][А-ЯЁA-Z-]+", token):
            patterns.append((unit.unit_id, re.compile(rf"(?<!\w){re.escape(token)}(?!\w)")))

    by_id = {c.clause_id: c for c in clauses}
    units_by_id = {u.unit_id: u for u in units}
    direct: dict[str, list[str]] = {}
    parent_ids = {c.parent_id for c in clauses if c.parent_id}
    titles: dict[str, list[Unit]] = {}
    for unit in units:
        if unit.kind == "unit":
            titles.setdefault(normalize(unit.name), []).append(unit)
    for clause in clauses:
        # A title equal to a sourced structural unit is a heading, not a duty
        # or an incidental mention. Only title clauses governing children qualify.
        title_matches = titles.get(normalize(_quote_name(clause.text).rstrip(":")), [])
        if clause.kind == "function" and clause.clause_id in parent_ids and len(title_matches) == 1:
            clause.kind = "heading"
            clause.unit_ids = list(dict.fromkeys([*clause.unit_ids, title_matches[0].unit_id]))
        direct[clause.clause_id] = list(dict.fromkeys(
            clause.unit_ids + [uid for uid, pat in patterns if pat.search(clause.text)]
        ))
    sibling_scopes = _sibling_role_scopes(clauses, direct, units_by_id)

    for clause in clauses:
        found = direct[clause.clause_id].copy()
        parent = by_id.get(clause.parent_id) if clause.parent_id else None
        while parent is not None:
            scoped = [
                uid for uid in direct[parent.clause_id]
                if uid in units_by_id and (
                    units_by_id[uid].kind == "role"
                    and any(c.clause_id == parent.clause_id and c.quote in parent.text
                            for c in units_by_id[uid].citations)
                    or _unit_scope(parent, units_by_id[uid])
                )
            ]
            if scoped:
                found.extend(uid for uid in scoped if uid not in found)
                break
            parent = by_id.get(parent.parent_id) if parent.parent_id else None
        found.extend(uid for uid in _scoped_sibling_roles(clause, by_id, sibling_scopes) if uid not in found)
        clause.unit_ids = found


def owner_keys(clause: Clause, clauses_by_id: dict[str, Clause], units_by_id: dict[str, Unit]) -> frozenset[str]:
    """Accountable role from a cited heading, never a delegate mentioned in prose."""
    current: Clause | None = clause
    while current is not None:
        roles = [
            unit_key(units_by_id[uid]) for uid in current.unit_ids
            if uid in units_by_id and units_by_id[uid].kind == "role"
            and any(c.clause_id == current.clause_id and c.quote in current.text
                    for c in units_by_id[uid].citations)
        ]
        if roles:
            return frozenset(roles)
        current = clauses_by_id.get(current.parent_id) if current.parent_id else None
    # A standalone, unnumbered role heading is a sibling of numbered duties.
    # Reconstruct its bounded source interval rather than treating mention
    # associations on a child as ownership evidence.
    scoped = _sibling_role_scopes(
        list(clauses_by_id.values()),
        {cid: [uid for uid in c.unit_ids if uid in units_by_id]
         for cid, c in clauses_by_id.items()},
        units_by_id,
    )
    sibling_roles = _scoped_sibling_roles(clause, clauses_by_id, scoped)
    if sibling_roles:
        return frozenset(unit_key(units_by_id[uid]) for uid in sibling_roles)

    current = clause
    while current is not None:
        keys = [
            unit_key(units_by_id[uid]) for uid in current.unit_ids
            if uid in units_by_id and _unit_scope(current, units_by_id[uid])
        ]
        if keys:
            return frozenset(keys)
        current = clauses_by_id.get(current.parent_id) if current.parent_id else None
    return frozenset()


def parse_document(document: Document, pages: list[dict]) -> ParseResult:
    """Clauses, units and parse warnings for one ingested document."""
    clauses, warnings, table_definitions = _parse_clauses(document.doc, pages)
    by_id = {clause.clause_id: clause for clause in clauses}
    governing_count = 0
    for clause in clauses:
        if clause.kind != "function":
            continue
        ancestors = []
        parent = by_id.get(clause.parent_id)
        while parent is not None:
            ancestors.append(parent.text)
            parent = by_id.get(parent.parent_id)
        if is_governance_statement(clause.text, ancestors):
            clause.kind = "other"
            governing_count += 1
    if governing_count:
        warnings.append(
            f"{document.doc}: {governing_count} governing/completeness statement(s) "
            "retained as source context, not counted as operational duties."
        )
    units = extract_units(document.doc, clauses, table_definitions=table_definitions)
    units_by_id = {unit.unit_id: unit for unit in units}
    unresolved_parents = [
        definition.clause.clause_id for definition in table_definitions
        if definition.parent_name and not any(
            units_by_id[uid].parent_unit_id
            for uid in definition.clause.unit_ids if uid in units_by_id
        )
    ]
    if unresolved_parents:
        shown = ", ".join(unresolved_parents[:5])
        warnings.append(
            f"{document.doc}: {len(unresolved_parents)} table parent reference(s) did "
            f"not uniquely resolve to a sourced unit; relationship not asserted ({shown})."
        )
    return ParseResult(clauses=clauses, units=units, warnings=warnings)


def parse_documents(documents: list[Document], pages_by_doc: dict[str, list[dict]]) -> ParseResult:
    """Parse every document; a document without pages yields a warning, not an exception."""
    result = ParseResult()
    for document in documents:
        pages = pages_by_doc.get(document.doc)
        if pages is None:
            result.warnings.append(f"{document.doc}: no ingested pages were supplied; document skipped.")
            continue
        parsed = parse_document(document, pages)
        result.clauses.extend(parsed.clauses)
        result.units.extend(parsed.units)
        result.warnings.extend(parsed.warnings)
    return result
