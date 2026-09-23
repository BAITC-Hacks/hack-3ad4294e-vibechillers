"""Regressions for observed Stage 3 integration defects; no quality/gold claims."""
from __future__ import annotations

import asyncio
from concurrent.futures import ThreadPoolExecutor
import hashlib
import json
from pathlib import Path
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from .. import db
from ..agent.audit_tools import AuditContext, create_audit_registry
from ..agent.loop import drive_audit
from ..agent.audit_llm import adjudicate_report
from ..api.audits import IngestedAudit, stream_audit
from ..config import Settings
from ..ingest import pipeline
from ..ingest.parsers import parse_any
from .models import ClauseRef, Finding
from .parser import owner_keys, parse_document
from .align import align_functions
from .report import resolve_alignment
from .test_context import audit
from .test_lineage import document, source, unit
from .lineage import analyze_domain


class SourceGuardTests(unittest.TestCase):
    def test_parallel_equal_filenames_keep_distinct_source_bytes(self):
        with tempfile.TemporaryDirectory() as folder:
            settings = Settings(_env_file=None, DATA_DIR=folder, DB_PATH=str(Path(folder)/'test.db'))
            barrier = threading.Barrier(2)
            original = pipeline.ingest_path
            def synchronized(path, **kwargs):
                barrier.wait(timeout=5)
                return original(path, **kwargs)
            payloads = [b'1. Alpha\n1.1. First duty.', b'1. Beta\n1.1. Second duty.']
            with patch.object(pipeline, 'get_settings', return_value=settings), patch.object(db, 'get_settings', return_value=settings):
                db.init_schema()
                with patch.object(pipeline, 'ingest_path', side_effect=synchronized), ThreadPoolExecutor(2) as pool:
                    futures = [pool.submit(pipeline.ingest_bytes, body, 'annex.txt') for body in payloads]
                    docs = [future.result(timeout=10) for future in futures]
            for body, doc in zip(payloads, docs):
                self.assertEqual(doc.doc_id, hashlib.sha256(body).hexdigest()[:16])
                self.assertEqual(''.join(p['text'] for p in doc.pages), body.decode())

    def test_structural_definitions_are_not_function_loss_predictions(self):
        clauses = [
            source('before','1','Отдел закупок'), source('after','2','Отдел снабжения'),
            source('before','1.1','Ведение реестра оборудования.','function','1',unit_ids=['1']),
            source('after','2.1','Ведение реестра оборудования.','function','2',unit_ids=['2']),
        ]
        units = [unit('before','1','Отдел закупок'),unit('after','2','Отдел снабжения')]
        findings = align_functions(clauses,units,['before'],['after'])
        self.assertEqual({(r.doc,r.clause_id) for f in findings for r in f.before+f.after},
                         {('before','1.1'),('after','2.1')})

    def test_duplicate_needs_distinct_successors_and_no_raw_rationale(self):
        report = audit('1. Функции\n1.1. Подготовка финансового отчета.', '1. Функции\n1.1. Подготовка технического отчета.')
        before=[ClauseRef(doc='before',clause_id='1.1')]; after=[ClauseRef(doc='after',clause_id='1.1')]
        pending=Finding(id='F1',status='unresolved',before=before,after=after,reason='review',method='exact',review_required=True)
        bad=resolve_alignment(pending,report.clauses,before=before,after=after*2,status='duplicate',reason='Два результата')
        self.assertEqual(bad.status,'unresolved')
        accepted=resolve_alignment(pending,report.clauses,before=before,after=after,status='changed',reason='Подтверждён штраф в миллион тенге.')
        self.assertEqual(accepted.status,'changed')
        self.assertNotIn('штраф',accepted.reason)

    def test_split_preserves_same_name_endpoint_and_negation_blocks_transition(self):
        docs=[document('before'),document('after')]
        units=[unit('before','1','Отдел закупок'),unit('after','2','Отдел закупок'),unit('after','3','Отдел логистики')]
        clauses=[source('before','1','Отдел закупок'),source('after','2','Отдел закупок'),source('after','3','Отдел логистики'),source('after','4','Отдел закупок разделён на Отдел закупок и Отдел логистики.','other')]
        changes,_,_=analyze_domain(docs,clauses,units,[])
        self.assertEqual([(c.status,len(c.before),len(c.after)) for c in changes],[('reorganised',1,2)])
        units=[unit('before','1','Отдел закупок'),unit('after','2','Отдел снабжения')]
        clauses=[source('before','1','Отдел закупок'),source('after','2','Отдел снабжения'),source('after','3','Отдел закупок не подлежит реорганизации в Отдел снабжения.','other')]
        changes,_,_=analyze_domain(docs,clauses,units,[])
        self.assertNotIn('reorganised',[c.status for c in changes])

    def test_participle_and_participation_do_not_assign_execution(self):
        docs=[document('before'),document('after')]
        units=[unit('before','1','Отдел контроля'),unit('after','1','Отдел контроля')]
        clauses=[source('before','1','Отдел контроля'),source('after','1','Отдел контроля'),source('after','1.1','Проверка подготовленных квартальных финансовых отчетов.','function','1',unit_ids=['1']),source('after','1.2','Проверка квартальных финансовых отчетов.','function','1',unit_ids=['1'])]
        changes,risks,_=analyze_domain(docs,clauses,units,[])
        self.assertEqual(changes[0].status,'retained')
        self.assertEqual(risks,[])
        units.append(unit('after','2','Отдел учёта')); clauses.append(source('after','2','Отдел учёта'))
        clauses.extend([source('after',f'{n}.3','Участие в подготовке квартальных финансовых отчетов.','function',str(n),unit_ids=[str(n)]) for n in (1,2)])
        self.assertEqual(analyze_domain(docs,clauses,units,[])[1],[])

    def test_specific_scope_conflict_is_not_overridden_by_common_broad_scope(self):
        docs=[document('before'),document('after')]
        units=[unit('after','1','Отдел закупок'),unit('after','2','Отдел снабжения')]
        clauses=[source('after','1','Отдел закупок'),source('after','2','Отдел снабжения')]
        for n,region in [(1,'Север'),(2,'Запад')]:
            clauses.append(source('after',f'{n}.1',f'Ведение реестра договоров в области закупок; в отношении региона {region}.','function',str(n),unit_ids=[str(n)]))
        self.assertEqual(analyze_domain(docs,clauses,units,[])[1],[])

    def test_merged_owner_column_does_not_import_another_columns_role(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            for merge,expected in [('B2:B3',{'директор'}),('A2:A3',{'отдел учета'})]:
                wb=Workbook(); ws=wb.active
                ws.append(['Подразделение','Должность','Функция'])
                ws.append(['Отдел учёта','Директор','Подготовка финансового отчета.'])
                ws.append(['Отдел учёта' if merge.startswith('B') else None,None,'Проверка финансового отчета.'])
                ws.merge_cells(merge); path=Path(folder)/'merged.xlsx'; wb.save(path)
                pages,_=parse_any(path); parsed=parse_document(document('after'),pages)
                clauses={c.clause_id:c for c in parsed.clauses}; units={u.unit_id:u for u in parsed.units}
                duty=next(c for c in parsed.clauses if c.text=='Проверка финансового отчета.')
                self.assertEqual(set(owner_keys(duty,clauses,units)),expected)
                self.assertEqual(duty.location.cell_range,'C3')

    def test_numbered_single_column_cells_preserve_coordinates_and_skip_formulas(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            wb = Workbook()
            ws = wb.active
            ws.title = 'Annex'
            ws['B1'] = '1. Обязанности'
            ws['B2'] = '1.1. Ведение реестра архивных дел.'
            ws['B3'] = '=B2'
            ws['B4'] = 'а) Проверка сроков хранения.'
            other = wb.create_sheet('Unheaded')
            other.append(['1. Обязанности', 'Подразделение'])
            other.append(['1.1. Ведение реестра.', 'Отдел учета'])
            path = Path(folder) / 'annex.xlsx'
            wb.save(path)
            pages, _ = parse_any(path)
            parsed = parse_document(document('after'), pages)
        clauses = {c.clause_id: c for c in parsed.clauses if c.location and c.location.sheet == 'Annex'}
        self.assertEqual(clauses['1.1'].text, 'Ведение реестра архивных дел.')
        self.assertEqual(clauses['1.1'].kind, 'function')
        self.assertEqual(clauses['1.1'].location.cell_range, 'B2')
        self.assertEqual(clauses['1.1/а'].parent_id, '1.1')
        self.assertEqual(clauses['1.1/а'].location.cell_range, 'B4')
        self.assertTrue(any(c.text == '=B2' and c.kind == 'other'
                            and c.location.cell_range == 'B3' for c in clauses.values()))
        self.assertFalse(any(c.kind == 'function' for c in parsed.clauses
                             if c.location and c.location.sheet == 'Unheaded'))

    def test_pdf_wrapped_numbered_and_lettered_paragraphs_keep_page_and_source_text(self):
        pages = [
            {'page': 1, 'physical_page': 1,
             'text': '1. Обязанности\r\n1.1. Ведение реестра\r\nархивных дел.\r\n'
                     'а) Проверка\r\nсроков хранения.\r\nНеозаглавленное приложение.'},
            {'page': 2, 'physical_page': 2,
             'text': '2. Контроль\r\n2.1. Составление отчета\r\nпо итогам сверки.'},
        ]
        parsed = parse_document(document('after'), pages)
        clauses = {c.clause_id: c for c in parsed.clauses}
        for clause_id, page_number in [('1.1', 1), ('1.1/а', 1), ('2.1', 2)]:
            clause = clauses[clause_id]
            self.assertEqual(clause.location.page, page_number)
            self.assertIn(clause.label + ' ' + clause.text, pages[page_number - 1]['text'])
        self.assertEqual(clauses['1.1'].text, 'Ведение реестра\r\nархивных дел.')
        self.assertEqual(clauses['1.1/а'].text, 'Проверка\r\nсроков хранения.')
        self.assertTrue(any(c.text == 'Неозаглавленное приложение.' and c.kind == 'other'
                            for c in parsed.clauses))


class AgentGuardTests(unittest.IsolatedAsyncioTestCase):
    async def test_same_response_read_and_finalize_is_not_completed(self):
        report=audit('1. Функции\n1.1. Ведение реестра.', '1. Функции\n1.1. Ведение реестра.')
        context=AuditContext.from_report(report)
        responses=iter([[('inspect_findings',{'finding_ids':[report.findings[0].id]}),('build_report',{'finding_ids':[],'conclusion':None})],[]])
        async def provider(*args,**kwargs):
            calls=[{'id':str(i),'type':'function','function':{'name':name,'arguments':json.dumps(params)}} for i,(name,params) in enumerate(next(responses))]
            yield {'done':{'tool_calls':calls}}
        settings=SimpleNamespace(llm_max_tokens=2048,llm_timeout_s=1)
        with patch('app.agent.loop.llm.stream',provider),patch('app.agent.loop.get_settings',return_value=settings):
            stats=await drive_audit('guard',run_id='guard',registry=create_audit_registry(context),system='guard',emit=None)
        self.assertFalse(stats.finalized)
        self.assertEqual(stats.invalid_calls,1)
        self.assertIsNone(context.last_report)

    async def test_idempotent_offer_and_abstention_allow_valid_finalization(self):
        report = audit('1. Функции\n1.1. Ведение реестра.', '1. Функции\n1.1. Ведение реестра.')
        pending = report.findings[0].model_copy(update={'status':'unresolved','method':'lexical','review_required':True})
        report = report.model_copy(update={'findings':[pending]})
        offer = {'finding_id':pending.id,'before':[r.model_dump() for r in pending.before],
                 'after':[r.model_dump() for r in pending.after]}
        responses = iter([
            [('inspect_findings',{'finding_ids':[pending.id]})],
            [('offer_candidates',offer),('resolve_alignment',{'finding_id':pending.id,'before':[],
                'after':[],'status':'unresolved','reason':'Недостаточно данных для решения.'})],
            [('build_report',{'finding_ids':[],'conclusion':None})],
        ])
        async def provider(*args,**kwargs):
            yield {'done':{'tool_calls':[{'id':str(i),'type':'function','function':{
                'name':name,'arguments':json.dumps(params)}} for i,(name,params) in enumerate(next(responses))]}}
        settings = SimpleNamespace(llm_configured=True,llm_model='guard',llm_max_tokens=2048,llm_timeout_s=1)
        with patch('app.agent.loop.llm.stream',provider),patch('app.agent.loop.get_settings',return_value=settings),patch('app.agent.audit_llm.get_settings',return_value=settings):
            result = await adjudicate_report(report)
        self.assertEqual(result.agent.status,'completed')
        self.assertEqual(result.findings[0].status,'unresolved')
        self.assertEqual(result.findings[0].before,pending.before)
        self.assertEqual(result.findings[0].after,pending.after)

    async def test_cancelled_tool_cannot_commit_after_deadline(self):
        report=audit('1. Функции\n1.1. Ведение реестра.', '1. Функции\n1.1. Ведение реестра.')
        context=AuditContext.from_report(report)
        entered=threading.Event(); release=threading.Event(); finished=threading.Event()
        def delayed(working,**kwargs):
            entered.set(); release.wait(timeout=5); working.warnings.append('late mutation'); finished.set(); return {}
        with patch.object(AuditContext,'list_findings',delayed):
            task=asyncio.create_task(context.invoke('list_findings'))
            self.assertTrue(await asyncio.to_thread(entered.wait,2))
            task.cancel()
            with self.assertRaises(asyncio.CancelledError): await task
            release.set(); self.assertTrue(await asyncio.to_thread(finished.wait,2))
        self.assertNotIn('late mutation',context.warnings)

    async def test_failed_final_storage_emits_error_not_success(self):
        with tempfile.TemporaryDirectory() as folder:
            settings=Settings(_env_file=None,DATA_DIR=folder,DB_PATH=str(Path(folder)/'test.db'))
            with patch.object(db,'get_settings',return_value=settings):
                db.init_schema(); db.create_run('persist',{'kind':'audit'})
                with db.session() as conn:
                    conn.execute("CREATE TRIGGER fail_final BEFORE INSERT ON agent_trace WHEN NEW.type='final' BEGIN SELECT RAISE(FAIL,'injected persistence failure'); END")
                documents=[document('before'),document('after')]
                input_=IngestedAudit(documents,{d.doc:[{'page':1,'text':'1. Функции\n1.1. Ведение реестра.'}] for d in documents},[])
                events=[e async for e in stream_audit('persist',input_,use_llm=False)]
                self.assertEqual([e.type for e in events if e.type in ('final','error')],['error'])
                self.assertFalse(any(e['type']=='final' for e in db.read_trace('persist')))


if __name__ == '__main__':
    unittest.main()
