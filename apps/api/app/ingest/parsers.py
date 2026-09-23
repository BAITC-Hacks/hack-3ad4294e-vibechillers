"""Document parsers: text plus source coordinates, without database access.

Every parser returns page dictionaries with ``page`` and ``text`` for existing
chunking callers. PDF ``physical_page``, DOCX ``block``, and XLSX ``sheet`` and
table ``rows`` additionally identify the original source; a Word block is not
a physical page. The audit parser uses row cells rather than flattening a
worksheet into supposed numbered duties.
"""

from __future__ import annotations

from pathlib import Path
from typing import Callable

# A page whose pypdfium2 text is shorter than this is treated as (likely)
# scanned/image-only and triggers the fallback chain.
_MIN_PDF_CHARS = 20

# Lazily constructed once per process; OCR model load is expensive.
_OCR_ENGINE = None


def _pdf_page_text_pypdfium2(page) -> str:
    """First-choice extraction for one open pypdfium2 page."""
    textpage = page.get_textpage()
    try:
        return textpage.get_text_range() or ""
    finally:
        textpage.close()


def _pdf_page_text_pdfplumber(pdf_path: Path, page_index: int) -> str:
    """Fallback extraction for a single page via pdfplumber."""
    import pdfplumber

    with pdfplumber.open(str(pdf_path)) as pdf:
        return (pdf.pages[page_index].extract_text() or "").strip()


def _pdf_page_text_ocr(page) -> str:
    """Optional last-resort text extraction from a rendered page.

    ``rapidocr-onnxruntime`` is an optional dependency. OCR text is not a
    structural interpretation of a diagram or a guarantee of completeness.
    """
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR

    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        _OCR_ENGINE = RapidOCR()

    bitmap = page.render(scale=2.0)
    image = bitmap.to_pil()
    result, _elapse = _OCR_ENGINE(np.array(image))
    if not result:
        return ""
    # rapidocr yields [box, text, confidence] triples in reading order.
    return "\n".join(str(line[1]) for line in result if line and line[1])


def parse_pdf(path: Path) -> list[dict]:
    """Extract pages in physical order, recording how each was read.

    Keep unreadable pages visible so a partially readable annex is not
    mistaken for a complete one. Ingestion rejects a wholly unreadable PDF.
    """
    import pypdfium2 as pdfium

    pages: list[dict] = []
    doc = pdfium.PdfDocument(str(path))
    try:
        for index in range(len(doc)):
            text = ""
            extraction = "unreadable"
            try:
                page = doc[index]
                text = _pdf_page_text_pypdfium2(page).strip()
                if text:
                    extraction = "pdfium"
                if len(text) < _MIN_PDF_CHARS:
                    try:
                        alt = _pdf_page_text_pdfplumber(Path(path), index)
                        if len(alt) > len(text):
                            text, extraction = alt, "pdfplumber"
                    except Exception:
                        pass
                if len(text) < _MIN_PDF_CHARS:
                    try:
                        alt = _pdf_page_text_ocr(page).strip()
                        if len(alt) > len(text):
                            text, extraction = alt, "ocr"
                    except ImportError:
                        pass
                    except Exception:
                        pass
                if text and len(text) < _MIN_PDF_CHARS:
                    extraction = "limited"
            except Exception:
                # Retain this physical page as an explicit limitation.
                pass
            pages.append({
                "page": index + 1,
                "physical_page": index + 1,
                "text": text or "[no extractable text]",
                "extraction": extraction,
            })
    finally:
        doc.close()
    return pages


def parse_xlsx(path: Path) -> list[dict]:
    """Extract each worksheet's original cells and physical coordinates.

    A merged cell is recorded only at its actual anchor. Other cells in the
    range remain empty, and the merge bounds are retained for conservative
    owner association by the audit parser.
    """
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=False, read_only=False)
    pages: list[dict] = []
    try:
        for sheet_index, ws in enumerate(wb.worksheets, start=1):
            rows: list[dict] = []
            for cells in ws.iter_rows():
                values = ["" if cell.value is None else str(cell.value) for cell in cells]
                if not any(value.strip() for value in values):
                    continue
                last = max(index for index, value in enumerate(values) if value.strip()) + 1
                values = values[:last]
                row_index = cells[0].row
                rows.append({
                    "row": row_index,
                    "cells": values,
                    "text": "\t".join(values),
                    "formulas": [
                        column for column, cell in enumerate(cells[:last], start=1)
                        if cell.data_type == "f"
                    ],
                })
            pages.append({
                "page": sheet_index,
                "sheet": ws.title,
                "text": "\n".join(row["text"] for row in rows),
                "rows": rows,
                "merges": [
                    {"min_row": span.min_row, "max_row": span.max_row,
                     "min_col": span.min_col, "max_col": span.max_col}
                    for span in ws.merged_cells.ranges
                ],
            })
    finally:
        wb.close()
    return pages


def parse_csv(path: Path) -> list[dict]:
    """Detect the encoding with charset_normalizer, read with pandas, render
    header + rows as tab-separated lines, all in a single page."""
    import pandas as pd
    from charset_normalizer import from_path

    best = from_path(str(path)).best()
    encoding = str(best.encoding) if best else "utf-8"
    frame = pd.read_csv(str(path), sep=",", encoding=encoding, dtype=str, keep_default_na=False)
    lines = ["\t".join(str(col) for col in frame.columns)]
    for row in frame.itertuples(index=False, name=None):
        lines.append("\t".join(str(value) for value in row))
    return [{"page": 1, "text": "\n".join(lines)}]


def parse_docx(path: Path) -> list[dict]:
    """Keep body-order paragraphs/tables and actual 1-based body block IDs.

    The legacy ``page`` slot stays sequential for chunking; ``block`` counts
    every body paragraph/table, including empty spacer paragraphs. Merged
    cells are never duplicated as if they contained separate source text.
    """
    import docx
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    pages: list[dict] = []
    block = 0
    for child in document.element.body:
        if isinstance(child, (CT_P, CT_Tbl)):
            block += 1
        if isinstance(child, CT_P):
            text = Paragraph(child, document).text
            if text.strip():
                pages.append({"page": len(pages) + 1, "block": block, "text": text})
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            rows: list[dict] = []
            anchors: dict[object, tuple[int, int]] = {}
            for row_index, row in enumerate(table.rows, start=1):
                values: list[str] = []
                merged: dict[int, tuple[int, int]] = {}
                for column, cell in enumerate(row.cells, start=1):
                    key = cell._tc
                    anchor = anchors.get(key)
                    if anchor is None:
                        anchors[key] = (row_index, column)
                        values.append(cell.text)
                    else:
                        values.append("")
                        merged[column] = anchor
                rows.append({
                    "row": row_index, "cells": values,
                    "text": "\t".join(values), "merged": merged,
                })
            pages.append({
                "page": len(pages) + 1,
                "block": block,
                "text": "\n".join(row["text"] for row in rows),
                "rows": rows,
            })
    if not pages:
        pages.append({"page": 1, "text": "", "block": 1})
    return pages


def parse_txt(path: Path) -> list[dict]:
    """Plain text, utf-8 with errors="replace", one page."""
    return [{"page": 1, "text": path.read_text(encoding="utf-8", errors="replace")}]


_PARSERS: dict[str, tuple[Callable[[Path], list[dict]], str]] = {
    ".pdf": (parse_pdf, "application/pdf"),
    ".docx": (
        parse_docx,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
    ".xlsx": (
        parse_xlsx,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ),
    ".csv": (parse_csv, "text/csv"),
    ".txt": (parse_txt, "text/plain"),
}


def parse_any(path: Path) -> tuple[list[dict], str]:
    """Dispatch on suffix; returns (pages, media_type). Raises ValueError for
    unsupported suffixes, naming the supported ones."""
    suffix = Path(path).suffix.lower()
    entry = _PARSERS.get(suffix)
    if suffix in {".doc", ".xls"}:
        replacement = ".docx" if suffix == ".doc" else ".xlsx"
        raise ValueError(
            f"Legacy {suffix} is not supported; convert it to {replacement} "
            "with a document editor and upload the converted file."
        )
    if entry is None:
        raise ValueError(
            f"Unsupported file suffix {suffix!r}; supported: "
            f"{', '.join(sorted(_PARSERS))}"
        )
    parser, media_type = entry
    return parser(Path(path)), media_type


def supported_suffixes() -> set[str]:
    """Exactly the suffixes ``parse_any`` accepts."""
    return set(_PARSERS)
